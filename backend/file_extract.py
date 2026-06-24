"""
Best-effort extraction of readable text from uploaded reference files, so
the AI team can use them as context. Runs entirely locally -- no upload,
no cloud parsing service, just local Python libraries.
"""
import csv
import io

MAX_CHARS_PER_FILE = 6000

IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "gif", "webp", "bmp"}

TEXT_EXTENSIONS = {
    "txt", "md", "markdown", "py", "js", "jsx", "ts", "tsx", "json", "csv",
    "html", "htm", "css", "yaml", "yml", "ini", "cfg", "log", "c", "h",
    "cpp", "hpp", "java", "go", "rb", "php", "sh", "bat", "ps1", "xml",
    "toml", "sql", "txt", "gitignore", "env",
}


def is_image(filename: str) -> bool:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    return ext in IMAGE_EXTENSIONS


def _truncate(text: str) -> str:
    if len(text) > MAX_CHARS_PER_FILE:
        return text[:MAX_CHARS_PER_FILE] + f"\n...[truncated, {len(text) - MAX_CHARS_PER_FILE} more characters]"
    return text


def extract_text(filename: str, raw: bytes):
    """Returns (text_or_None, note_or_None)."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext in TEXT_EXTENSIONS:
        try:
            try:
                text = raw.decode("utf-8")
            except UnicodeDecodeError:
                text = raw.decode("latin-1")
            return _truncate(text), None
        except Exception as e:
            return None, f"text decode failed: {e}"

    if ext == "pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(raw))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
            if not text.strip():
                return None, "PDF appears to be scanned/image-based; no extractable text"
            return _truncate(text), None
        except Exception as e:
            return None, f"PDF parsing failed: {e}"

    if ext == "docx":
        try:
            import docx
            d = docx.Document(io.BytesIO(raw))
            text = "\n".join(p.text for p in d.paragraphs)
            return _truncate(text), None
        except Exception as e:
            return None, f"Word document parsing failed: {e}"

    if ext in ("xlsx", "xlsm"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
            out = io.StringIO()
            for sheet in wb.worksheets:
                out.write(f"--- Sheet: {sheet.title} ---\n")
                writer = csv.writer(out)
                for i, row in enumerate(sheet.iter_rows(values_only=True)):
                    if i > 500:
                        out.write("...[truncated, sheet has more rows]\n")
                        break
                    writer.writerow(["" if v is None else v for v in row])
            return _truncate(out.getvalue()), None
        except Exception as e:
            return None, f"Excel parsing failed: {e}"

    return None, "binary/unsupported file type"
