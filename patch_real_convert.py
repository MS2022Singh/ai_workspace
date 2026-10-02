import re
from pathlib import Path

main_path = Path("main.py")
src = main_path.read_text(encoding="utf-8")

# Match the entire /api/convert block
pattern = re.compile(
    r'@app\.post\("/api/convert"\).*?(?=@app\.post\(|@app\.get\()',
    re.DOTALL
)

new_convert = r'''@app.post("/api/convert")
async def convert_file(file: UploadFile = File(...), target_format: Optional[str] = Form("pdf")):
    """
    Real file conversion endpoint.
    Supports: txt/md -> pdf, docx, html; images -> png/jpg; csv -> xlsx
    """
    import uuid, os
    from pathlib import Path as _Path

    try:
        # Save uploaded file to a temp working directory
        work = _Path("storage_outputs")
        work.mkdir(exist_ok=True)
        stem = _Path(file.filename).stem or "uploaded"
        src_ext = _Path(file.filename).suffix.lstrip(".").lower() or "bin"
        content = await file.read()

        tmp_src = work / f"src_{uuid.uuid4().hex[:8]}_{stem}.{src_ext}"
        tmp_src.write_bytes(content)

        target = (target_format or "pdf").lower().lstrip(".")
        out_name = f"{stem}_converted.{target}"
        out_path = work / out_name

        # -------- TEXT-LIKE -> PDF --------
        if target == "pdf" and src_ext in ("txt", "md", "log", "csv", "json", "html"):
            text = content.decode("utf-8", errors="replace")
            from fpdf import FPDF
            pdf = FPDF()
            pdf.add_page()
            font_path = _Path("assets/fonts/DejaVuSans.ttf")
            if font_path.exists():
                pdf.add_font("DejaVu", "", str(font_path), uni=True)
                pdf.set_font("DejaVu", "", 11)
            else:
                pdf.set_font("Helvetica", "", 11)
            for para in text.split("\n"):
                try:
                    pdf.multi_cell(0, 6, para if para else " ")
                except Exception:
                    pdf.multi_cell(0, 6, para.encode("latin-1", "replace").decode("latin-1"))
            pdf.output(str(out_path))

        # -------- TEXT-LIKE -> DOCX --------
        elif target == "docx" and src_ext in ("txt", "md", "log", "csv", "json"):
            text = content.decode("utf-8", errors="replace")
            from docx import Document
            doc = Document()
            for para in text.split("\n"):
                doc.add_paragraph(para)
            doc.save(str(out_path))

        # -------- TEXT-LIKE -> HTML --------
        elif target == "html" and src_ext in ("txt", "md", "log"):
            text = content.decode("utf-8", errors="replace")
            html = "<!DOCTYPE html><html><head><meta charset='utf-8'><title>" + stem + "</title></head><body><pre>" + text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;") + "</pre></body></html>"
            out_path.write_text(html, encoding="utf-8")

        # -------- IMAGE -> IMAGE --------
        elif src_ext in ("png", "jpg", "jpeg", "webp", "bmp", "gif", "tiff") and target in ("png", "jpg", "jpeg", "webp", "bmp", "tiff"):
            from PIL import Image
            img = Image.open(tmp_src).convert("RGB" if target in ("jpg", "jpeg") else "RGBA")
            save_target = "JPEG" if target in ("jpg", "jpeg") else target.upper()
            if save_target == "JPG": save_target = "JPEG"
            img.save(str(out_path), save_target)

        # -------- CSV -> XLSX --------
        elif target == "xlsx" and src_ext == "csv":
            import csv
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            with open(tmp_src, "r", encoding="utf-8", errors="replace") as f:
                for row in csv.reader(f):
                    ws.append(row)
            wb.save(str(out_path))

        else:
            return {
                "status": "unsupported",
                "detail": f"Conversion {src_ext} -> {target} not implemented",
                "response": f"Conversion {src_ext.upper()} -> {target.upper()} is not supported yet.",
                "result": "unsupported"
            }

        # Cleanup temp
        try: tmp_src.unlink()
        except Exception: pass

        return {
            "status": "success",
            "source_file": file.filename,
            "target_format": target.upper(),
            "output_file": out_name,
            "size_bytes": out_path.stat().st_size,
            "download_url": f"/api/output/download/{out_name}",
            "response": f"Converted '{file.filename}' to {target.upper()} ({out_path.stat().st_size} bytes).",
            "result": "Conversion complete"
        }

    except Exception as e:
        import traceback
        return {
            "status": "error",
            "detail": str(e),
            "trace": traceback.format_exc()[-500:],
            "response": f"Conversion failed: {e}",
            "result": "error"
        }

'''

new_src, count = pattern.subn(new_convert, src, count=1)
if count > 0:
    main_path.write_text(new_src, encoding="utf-8")
    print(f"PATCHED: /api/convert replaced with real converter")
else:
    print("ERROR: could not locate /api/convert block")
